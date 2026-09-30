import pdfplumber
import os

# DOCX
from docx import Document

# OCR image
from ..config import PDF_PAGE_LIMIT, PYTESSERACT_CMD
import pytesseract
from PIL import Image

# Excel
from openpyxl import load_workbook
pytesseract.pytesseract.tesseract_cmd = PYTESSERACT_CMD

from odf.opendocument import load
from odf.text import P as Ptext

# Function to recursively extract text from ODT elements
def extract_text(element):
    text = ""

    for node in element.childNodes:
        if node.nodeType == node.TEXT_NODE:
            text += node.data
        else:
            text += extract_text(node)

    return text

# Function to read ODT files
def read_odt(path):
    try:
        doc = load(path)
        paragraphs = doc.getElementsByType(Ptext)

        text = "\n".join(extract_text(p) for p in paragraphs)

        return text.strip()

    except Exception as e:
        print(f"Erreur lecture ODT {path}: {e}")
        return ""
        
# Function to read content from various file types
def read_file_content(path):
    ext = os.path.splitext(path)[1].lower()

    try:
        # -------------------------
        # PDF
        # -------------------------
        if ext == ".pdf":
            text = ""

            try:
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages[:PDF_PAGE_LIMIT]:  # limite pour perf
                        text += page.extract_text() or ""
            except Exception as error:
                print(f"Erreur d'extraction PDF {path} : {error}")
                text = ""

            # fallback OCR si vide
            if not text.strip() and ext == ".pdf":
                try:
                    from pdf2image import convert_from_path
                    images = convert_from_path(path, first_page=1, last_page=2)

                    for img in images:
                        text += pytesseract.image_to_string(img)
                except Exception as error:
                    print(f"Erreur OCR du PDF {path} : {error}")

            return text.strip()

        # -------------------------
        # DOCX
        # -------------------------
        elif ext == ".docx":
            try:
                doc = Document(path)
                return "\n".join(p.text for p in doc.paragraphs)
            except Exception as error:
                print(f"Erreur de lecture DOCX {path} : {error}")
                return ""
        
        # -------------------------
        # ODT
        # -------------------------
        elif ext == ".odt":
            return read_odt(path)

        # -------------------------
        # TXT / MD
        # -------------------------
        elif ext in [".txt", ".md", ".csv"]:
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            except Exception as error:
                print(f"Erreur de lecture du fichier texte {path} : {error}")
                return ""

        # -------------------------
        # Excel
        # -------------------------
        elif ext == ".xlsx":
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                text = ""

                for sheet in wb:
                    for row in sheet.iter_rows(values_only=True):
                        text += " ".join([str(cell) for cell in row if cell]) + "\n"

                return text
            except Exception as error:
                print(f"Erreur de lecture XLSX {path} : {error}")
                return ""

        # -------------------------
        # Images (OCR)
        # -------------------------
        elif ext in [".png", ".jpg", ".jpeg", ".bmp"]:
            try:
                return pytesseract.image_to_string(Image.open(path))
            except Exception as error:
                print(f"Erreur OCR de l'image {path} : {error}")
                return ""

        # -------------------------
        # Fallback
        # -------------------------
        else:
            return ""

    except Exception as e:
        print(f"Erreur lecture fichier {path}: {e}")
        return ""
